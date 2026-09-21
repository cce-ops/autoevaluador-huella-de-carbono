"""
Comprobaciones DETERMINISTAS (sin IA) sobre el Excel "Datos tareas HC 2026v3".

Idea: lo que se puede medir con código no se le pregunta al LLM. Estos hallazgos
se inyectan en el prompt como hechos y limitan la nota máxima (ver app.py).

Todas las celdas se citan con su coordenada (p. ej. '1-INVENTARIO!K16') para que
el feedback sea verificable.
"""
import io
import re
import openpyxl

# ----------------------------------------------------------------------------
# CONFIGURACIÓN (revisar por el profesor)
# ----------------------------------------------------------------------------
HOJA_INV = "1-INVENTARIO"
HOJA_CALC = "2-CÁLCULO"
FILAS_INV = range(14, 27)          # A14:P26 según la guía
COLS_DATOS = "BCEFGHIJKLMNOP"      # columnas numéricas del inventario (D = texto: tipo de combustible)
ALCANCES_ESPERADOS = {6: 3, 7: 2, 8: 1, 9: 3, 10: 3}   # agua, elec., comb., resid., transp.

# Palabras que delatan datos de ejemplo/estadísticos en textos libres
MARCADORES = ["prueba", "test", "ejemplo", "media", "promedio", "estadíst", "estadist",
              "estimad", "aprox", "típic", "tipic", "genéric", "generic", "lorem"]

# Valores de referencia para detectar "media estadística copiada".
# ⚠️ [No verificado] Los pone/confirma el profesor con su fuente (INE, ARC, etc.).
# Si un total del alumno cae dentro de la tolerancia, se marca como SOSPECHOSO
# (no como prueba): es una alerta para que el LLM pida evidencia.
REFERENCIAS_MEDIAS = {
    # "clave": (valor, unidad, tolerancia relativa, descripción)
    "residuos_total": (470.0, "kg/año", 0.03, "Total residuos municipales por persona (verificar fuente)"),
}


def _num(v):
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def _es_redondo(x, paso):
    return abs(x / paso - round(x / paso)) < 1e-9


def cargar(file_bytes):
    """Abre el libro con valores cacheados (data_only) y con fórmulas."""
    bio = io.BytesIO(file_bytes)
    wb_v = openpyxl.load_workbook(bio, data_only=True)
    bio.seek(0)
    wb_f = openpyxl.load_workbook(bio, data_only=False)
    return wb_v, wb_f


def volcado_con_coordenadas(file_bytes, max_celdas=4000):
    """Texto para el LLM: solo celdas no vacías, con coordenada. Sustituye a df.to_string()."""
    wb_v, _ = cargar(file_bytes)
    out = []
    for ws in wb_v:
        out.append(f"\n--- HOJA: {ws.title} ---")
        n = 0
        for row in ws.iter_rows():
            for c in row:
                if c.value is not None and str(c.value).strip() != "":
                    out.append(f"{ws.title}!{c.coordinate} = {c.value!r}")
                    n += 1
                    if n >= max_celdas:
                        out.append("[... hoja truncada ...]")
                        break
            if n >= max_celdas:
                break
    return "\n".join(out)


def analizar_excel(file_bytes):
    """Devuelve dict con 'hallazgos' (lista de dicts) y 'metricas'."""
    h = []   # hallazgos: {"gravedad": "alta|media|baja", "codigo": str, "texto": str}
    m = {}

    def add(g, cod, txt):
        h.append({"gravedad": g, "codigo": cod, "texto": txt})

    try:
        wb_v, wb_f = cargar(file_bytes)
    except Exception as e:
        add("alta", "ARCHIVO_ILEGIBLE", f"No se pudo abrir el Excel: {e}")
        return {"hallazgos": h, "metricas": m}

    if HOJA_INV not in wb_v.sheetnames or HOJA_CALC not in wb_v.sheetnames:
        add("media", "HOJAS_AUSENTES",
            f"No se encuentran las hojas '{HOJA_INV}' y/o '{HOJA_CALC}'. "
            f"Hojas presentes: {wb_v.sheetnames}. (Puede no ser el Excel de la actividad.)")
        return {"hallazgos": h, "metricas": m}

    inv = wb_v[HOJA_INV]
    calc = wb_v[HOJA_CALC]

    # ---- 1) Marcadores de ejemplo/estadística en textos ------------------------------------
    textos = []
    for r in list(range(2, 11)) + list(FILAS_INV):
        v = inv[f"A{r}"].value
        if isinstance(v, str) and v.strip():
            textos.append((f"{HOJA_INV}!A{r}", v))
    for coord, t in textos:
        for mk in MARCADORES:
            if mk in t.lower():
                add("alta", "MARCADOR_TEXTO", f"{coord} contiene «{mk}»: «{t}».")
                break

    # ---- 2) Redondez de los datos numéricos -------------------------------------------------
    nums = []
    for r in FILAS_INV:
        for col in COLS_DATOS:
            v = inv[f"{col}{r}"].value
            if _num(v) and v != 0:
                nums.append((f"{HOJA_INV}!{col}{r}", float(v)))
    m["n_valores_inventario"] = len(nums)
    if nums:
        r5 = sum(_es_redondo(v, 5) for _, v in nums) / len(nums)
        r10 = sum(_es_redondo(v, 10) for _, v in nums) / len(nums)
        enteros = sum(_es_redondo(v, 1) for _, v in nums) / len(nums)
        m.update({"pct_multiplo_5": round(r5, 2), "pct_multiplo_10": round(r10, 2),
                  "pct_enteros": round(enteros, 2)})
        if len(nums) >= 6 and r5 >= 0.8:
            add("alta", "DATOS_REDONDOS",
                f"{r5:.0%} de los {len(nums)} valores del inventario son múltiplos de 5 "
                f"({r10:.0%} de 10; {enteros:.0%} enteros). Datos medidos/pesados suelen ser irregulares "
                f"(p. ej. pesaje × 52,14). Patrón compatible con estimaciones genéricas.")
        elif len(nums) >= 6 and r5 >= 0.6:
            add("media", "DATOS_CASI_REDONDOS",
                f"{r5:.0%} de los valores son múltiplos de 5. Pedir evidencia.")
    else:
        add("alta", "INVENTARIO_VACIO", f"No hay valores numéricos en {HOJA_INV}!A14:P26.")

    # ---- 3) Residuos: ¿pesaje real? ---------------------------------------------------------
    # Fracciones F..J: papel, envases, vidrio, orgánico, resto
    fr = {}
    for col, nombre in zip("FGHIJ", ["papel", "envases", "vidrio", "organico", "resto"]):
        fr[nombre] = sum(inv[f"{col}{r}"].value or 0 for r in FILAS_INV if _num(inv[f"{col}{r}"].value))
    tot = sum(fr.values())
    m["residuos_kg"] = {k: round(v, 2) for k, v in fr.items()}
    m["residuos_total_kg"] = round(tot, 2)
    if tot > 0:
        # Un pesaje de 7 días extrapolado (× 52,14) produce múltiplos de ~0,05*52,14: casi nunca enteros redondos.
        if all(_es_redondo(v, 5) for v in fr.values() if v):
            add("alta", "RESIDUOS_REDONDOS",
                f"Todas las fracciones de residuos son múltiplos de 5 kg {fr}: no es el resultado típico "
                f"de un pesaje de 3–7 días extrapolado al año.")
        ref = REFERENCIAS_MEDIAS.get("residuos_total")
        if ref and abs(tot - ref[0]) <= ref[2] * ref[0]:
            add("alta", "COINCIDE_CON_MEDIA",
                f"El total de residuos ({tot:.0f} kg/año) está a ±{ref[2]:.0%} del valor de referencia "
                f"{ref[0]:.0f} {ref[1]} ({ref[3]}).")

    # ---- 4) Coherencia 2-CÁLCULO -----------------------------------------------------------
    suma = 0
    for r, esperado in ALCANCES_ESPERADOS.items():
        a = calc[f"A{r}"].value
        c = calc[f"C{r}"].value
        if a != esperado:
            add("media", "ALCANCE", f"{HOJA_CALC}!A{r} = {a!r}; la actividad exige {esperado}.")
        if not _num(c):
            add("alta", "CALC_VACIO", f"{HOJA_CALC}!C{r} vacío o no numérico ({c!r}).")
        else:
            suma += c
            if c > 50:
                add("media", "UNIDAD_KG_VS_T",
                    f"{HOJA_CALC}!C{r} = {c}: valor muy alto para t CO₂eq/persona; ¿kg en vez de t?")
    c11 = calc["C11"].value
    if _num(c11) and abs(c11 - suma) > 1e-6:
        add("media", "TOTAL_INCOHERENTE", f"{HOJA_CALC}!C11 = {c11} pero la suma de C6:C10 es {suma:.4f}.")
    m["huella_total_t"] = round(suma, 4)

    # ---- 5) Signo de la mejora en escenarios (bug de plantilla original) --------------------
    for nombre in ("ESCENARIO A", "ESCENARIO B"):
        if nombre in wb_f.sheetnames:
            f = wb_f[nombre]["H25"].value
            if isinstance(f, str) and re.sub(r"\s", "", f).upper() == "=G25-C25":
                add("baja", "MEJORA_SIGNO",
                    f"{nombre}!H25 usa '=G25-C25': una reducción sale NEGATIVA. La guía indica '=C25-G25' "
                    f"y I25 '=SI(C25=0;0;H25/C25)'. Es un fallo de la plantilla, no penalizar al alumno.")
            # Escenario sin datos propios
            g = [wb_v[nombre][f"G{r}"].value for r in range(25, 30)]
            if not all(_num(x) for x in g):
                add("alta", "ESC_SIN_RESULTADOS", f"{nombre}!G25:G29 incompleto: {g}.")

    # ---- 6) Escenarios: ¿cambian entradas físicas? -----------------------------------------
    for nombre in ("ESCENARIO A", "ESCENARIO B"):
        if nombre in wb_v.sheetnames:
            ws = wb_v[nombre]
            n = sum(1 for r in range(10, 20) for col in COLS_DATOS if _num(ws[f"{col}{r}"].value))
            if n == 0:
                add("alta", "ESC_SIN_INVENTARIO", f"{nombre}!A10:P19 no contiene datos de actividad.")

    return {"hallazgos": h, "metricas": m}


def evidencias_presentes(nombres_archivos):
    """¿Ha subido el alumno el documento de inventario con evidencias y las OCCC?"""
    n = [x.lower() for x in nombres_archivos]
    return {
        "doc_inventario": any(("inventario" in x and x.endswith((".pdf", ".docx"))) for x in n),
        "occc": any(x.endswith((".xlsm",)) or "occc" in x for x in n),
    }
